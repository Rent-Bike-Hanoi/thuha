#!/usr/bin/env python3
"""Export the SEO-recovery master matrix workbook (13 sheets) for thuha.rentbikehanoi.com.

Reads the data files in reports/seo-recovery/data/ (CSV/TSV/JSONL/JSON, utf-8) plus
data/keyword-metrics.jsonl and writes one .xlsx with exactly these sheets, in order:

  01_DASHBOARD          live COUNTIF/COUNTIFS/SUM formulas over the other sheets
  02_ALL_URL_INVENTORY  every rendered URL (+ REMOVED_PRIOR rows appended)
  03_INDEX_RECOVERY     per-URL recovery tracker (editable fix_method / noindex_after / notes)
  04_KEYWORD_MATRIX     one row per indexable URL + one row per GSC query candidate
  05_DUPLICATE_CLUSTERS similarity clusters / near-duplicates
  06_CONTENT_AUDIT      on-page audit flags per indexable or rewrite-candidate URL
  07_TECHNICAL_SEO      site-wide findings block + per-URL technical issues block
  08_INTERNAL_LINKS     every internal link edge with BROKEN / LINK_TO_NOINDEX errors
  09_CHANGELOG          changelog.jsonl
  10_QA_REPORT          qa_report.jsonl
  11_CHECKPOINT         checkpoint.json as key/value rows
  12_GSC_PERFORMANCE    daily / pages / inspection / query blocks
  13_REPOSITORY_AUDIT   repo_audit.json

Usage:
    python3 scripts/seo_matrix_export.py [--data DIR] [--out FILE] [--verify]

Pure stdlib + openpyxl.  Idempotent: the output file is overwritten.  Optional inputs
(changelog.jsonl, qa_report.jsonl, checkpoint.json, technical_findings.json,
repo_audit.json, keyword_matrix.csv) are written as header-only tables when absent.
"""
import argparse
import collections
import csv
import datetime
import json
import re
import sys
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATA = ROOT / 'reports' / 'seo-recovery' / 'data'
DEFAULT_OUT = ROOT / 'reports' / 'seo-recovery' / 'THUHA-SEO-MASTER-MATRIX-2026-10-30.xlsx'
KW_METRICS = ROOT / 'data' / 'keyword-metrics.jsonl'
SITE = 'https://thuha.rentbikehanoi.com'
GSC_NOTE = ('Source: Google Search Console via Windsor.ai connector + URL Inspection API (OpenSEO); '
            'data fetched 2026-10-30; INDEX_ELIGIBLE != INDEXED')

SHEET_NAMES = ['01_DASHBOARD', '02_ALL_URL_INVENTORY', '03_INDEX_RECOVERY', '04_KEYWORD_MATRIX',
               '05_DUPLICATE_CLUSTERS', '06_CONTENT_AUDIT', '07_TECHNICAL_SEO', '08_INTERNAL_LINKS',
               '09_CHANGELOG', '10_QA_REPORT', '11_CHECKPOINT', '12_GSC_PERFORMANCE', '13_REPOSITORY_AUDIT']
STATUS_VALUES = ['OPEN', 'IN_PROGRESS', 'DONE', 'BLOCKED', 'NEEDS_REVIEW', 'DONE_BY_RULE']
PRIORITY_VALUES = ['P0', 'P1', 'P2', 'P3', 'P4']
CATEGORIES = ['INDEXABLE', 'A', 'B', 'C', 'D', 'E']
PAGE_TYPES = ['home', 'pillar', 'manual_article', 'static_page', 'hub', 'pagination', 'author',
              'article_v1', 'article_v2']
NEEDS_SOURCES = {'xe-may', 'xe-dien', 'luat-giao-thong', 'bao-duong', 'du-lich'}
ARTICLE_TYPES = {'article_v1', 'article_v2', 'manual_article', 'pillar'}
RECOVERY_CATEGORIES = {'A', 'B', 'C', 'E'}

# Columns coerced to int/float when the CSV value looks numeric.
NUMERIC_COLS = {'title_len', 'title_full_len', 'meta_desc_len', 'h1_count', 'h2_count', 'h3_count',
                'word_count', 'src_wordCount', 'sections', 'body_link_count', 'external_links',
                'img_count', 'img_missing_alt', 'sources_count', 'bytes', 'inlinks_body',
                'inlinks_related', 'inlinks_total', 'factoryVersion', 'gsc_clicks', 'gsc_impressions',
                'gsc_position', 'max_sim', 'max_sim_indexable', 'cluster_size', 'n_over_80',
                'n_over_60', 'clicks', 'impressions', 'ctr', 'position', 'score', 'volume', 'kd',
                'affected_count', 'qa_score', 'Volume', 'KD', 'GSC impressions', 'GSC clicks',
                'GSC position', 'Internal links in (inlinks_total)'}
# Free-text columns that get wrap_text (everything else: wrap off, width capped at 60).
WRAP_COLS = {'note', 'notes', 'Notes', 'reason', 'evidence', 'fix', 'finding', 'before', 'after',
             'checks', 'critical', 'differentiation_plan', 'value', 'Note'}
MAX_CELL = 32000  # Excel limit is 32,767 characters per cell

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill('solid', fgColor='FFD9E1F2')
SECTION_FILL = PatternFill('solid', fgColor='FFEDEDED')
FILL_RED = PatternFill('solid', fgColor='FFFFC7CE')
FILL_GREEN = PatternFill('solid', fgColor='FFC6EFCE')
FILL_YELLOW = PatternFill('solid', fgColor='FFFFEB9C')
WRAP = Alignment(wrap_text=True, vertical='top')


# ----------------------------------------------------------------------------- readers
def read_csv(path, delimiter=','):
    """Return (fieldnames, rows-as-dicts); ([], []) when the file is missing."""
    path = Path(path)
    if not path.exists():
        return [], []
    csv.field_size_limit(1 << 30)
    with open(path, newline='', encoding='utf-8') as fh:
        rd = csv.DictReader(fh, delimiter=delimiter)
        rows = list(rd)
        return list(rd.fieldnames or []), rows


def read_jsonl(path):
    out = []
    path = Path(path)
    if not path.exists():
        return out
    with open(path, encoding='utf-8') as fh:
        for ln, line in enumerate(fh, 1):
            line = line.strip()
            if not line:
                continue
            try:
                out.append(json.loads(line))
            except json.JSONDecodeError as exc:
                print(f'WARN {path.name}:{ln}: skipped malformed JSON ({exc})', file=sys.stderr)
    return out


def read_json(path, default=None):
    path = Path(path)
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding='utf-8'))
    except json.JSONDecodeError as exc:
        print(f'WARN {path.name}: unreadable JSON ({exc}); treated as absent', file=sys.stderr)
        return default


def find_file(data, name, prefix=None, suffix=''):
    """Exact file name first; otherwise the newest file matching prefix*suffix."""
    p = data / name
    if p.exists():
        return p
    if prefix:
        cands = sorted(data.glob(f'{prefix}*{suffix}'))
        if cands:
            return cands[-1]
    return p


# ----------------------------------------------------------------------------- value helpers
INT_RE = re.compile(r'^-?\d+$')
FLOAT_RE = re.compile(r'^-?(\d+\.\d*|\.\d+)([eE][-+]?\d+)?$')


def clean(v):
    """Make any value safe for an openpyxl cell ('' -> None, bools -> 'True'/'False')."""
    if v is None:
        return None
    if isinstance(v, bool):
        return 'True' if v else 'False'
    if isinstance(v, (int, float)):
        return v
    if not isinstance(v, str):
        v = json.dumps(v, ensure_ascii=False)
    if v == '':
        return None
    v = ILLEGAL_CHARACTERS_RE.sub('', v)
    if len(v) > MAX_CELL:
        v = v[:MAX_CELL] + ' …[truncated]'
    if v.startswith('='):          # keep openpyxl from storing data as a formula
        v = ' ' + v
    return v


def num(v):
    """Coerce numeric-looking strings to int/float; '' -> None; other text unchanged."""
    if v is None or (isinstance(v, (int, float)) and not isinstance(v, bool)):
        return v
    s = str(v).strip()
    if s == '':
        return None
    if INT_RE.match(s):
        return int(s)
    if FLOAT_RE.match(s):
        return float(s)
    return clean(s)


def conv(col, v):
    return num(v) if col in NUMERIC_COLS else clean(v)


def to_f(v, default=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def to_i(v, default=0):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


def joined(v, sep='; '):
    if v is None or v == '':
        return None
    if isinstance(v, (list, tuple)):
        return clean(sep.join(x if isinstance(x, str) else json.dumps(x, ensure_ascii=False) for x in v))
    if isinstance(v, dict):
        return clean(sep.join(f'{k}={x}' for k, x in v.items()))
    return clean(v)


def path_of(u):
    """'/a/b/' from either a path or an absolute URL on the site."""
    u = (u or '').strip()
    if u.startswith('http'):
        u = re.sub(r'^https?://[^/]+', '', u) or '/'
    return u


def is_true(v):
    return str(v).strip().lower() in ('true', '1', 'yes')


# ----------------------------------------------------------------------------- sheet writer
class Sheet:
    """Thin wrapper over a worksheet that tracks the current row and column widths."""

    def __init__(self, wb, name):
        self.ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        self.row = 0
        self.widths = collections.defaultdict(int)

    def append(self, values):
        self.ws.append(values)
        self.row += 1
        if self.row <= 600:
            for i, v in enumerate(values, 1):
                if v is not None:
                    self.widths[i] = max(self.widths[i], min(80, len(str(v))))
        return self.row

    def blank(self, n=1):
        for _ in range(n):
            self.append([])

    def title(self, text):
        r = self.append([text])
        self.ws.cell(r, 1).font = Font(bold=True, size=12)
        return r

    def header(self, headers):
        r = self.append(list(headers))
        for c in range(1, len(headers) + 1):
            cell = self.ws.cell(r, c)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
        return r

    def table(self, headers, rows, title=None):
        """Write [title], header and rows. Returns (header_row, first_data_row, last_row)."""
        if title:
            self.title(title)
        h = self.header(headers)
        for r in rows:
            self.append(r)
        return h, h + 1, self.row

    def col(self, headers, name):
        return get_column_letter(list(headers).index(name) + 1)

    def finish(self, headers, header_row, last_row, freeze='A2', autofilter=True,
               wrap_cols=WRAP_COLS, widths=None):
        ws = self.ws
        headers = list(headers)
        for i, h in enumerate(headers, 1):
            w = max(self.widths.get(i, 0), len(str(h)))
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, w + 2))
        for letter, w in (widths or {}).items():
            ws.column_dimensions[letter].width = w
        if freeze:
            ws.freeze_panes = freeze
        if autofilter:
            ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{max(last_row, header_row)}'
        for name in wrap_cols:
            if name in headers:
                c = headers.index(name) + 1
                for r in range(header_row + 1, last_row + 1):
                    ws.cell(r, c).alignment = WRAP

    def status_cf(self, col_letter, first, last):
        if last < first:
            return
        rng = f'{col_letter}{first}:{col_letter}{last}'
        for values, fill in ((('BLOCKED', 'FAIL'), FILL_RED),
                             (('DONE', 'PASS', 'DONE_BY_RULE'), FILL_GREEN),
                             (('NEEDS_REVIEW', 'IN_PROGRESS'), FILL_YELLOW)):
            for v in values:
                self.ws.conditional_formatting.add(
                    rng, CellIsRule(operator='equal', formula=[f'"{v}"'], fill=fill))

    def list_dv(self, col_letter, first, last, values):
        if last < first:
            return
        dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
        dv.errorStyle = 'warning'
        dv.errorTitle = 'Value not in list'
        dv.error = 'Pick a value from the list, or confirm to keep the typed text.'
        self.ws.add_data_validation(dv)
        dv.add(f'{col_letter}{first}:{col_letter}{last}')

    def number_format(self, col_letter, first, last, fmt):
        for r in range(first, last + 1):
            self.ws[f'{col_letter}{r}'].number_format = fmt


# ----------------------------------------------------------------------------- data loading
def load_all(data):
    d = {}
    d['inv_fields'], d['inv'] = read_csv(data / 'inventory.csv')
    _, d['removed'] = read_csv(find_file(data, 'removed_prior_2026-10-10.csv', 'removed_prior_', '.csv'))
    _, d['cls'] = read_csv(data / 'classification.csv')
    _, d['sim'] = read_csv(data / 'similarity.csv')
    d['links_fields'], d['links'] = read_csv(data / 'internal_links.csv')
    _, d['gsc_pages'] = read_csv(find_file(data, 'gsc_pages_2026-08-01_10-27.tsv', 'gsc_pages_', '.tsv'), '\t')
    _, d['gsc_daily'] = read_csv(find_file(data, 'gsc_daily_2026-10-05_10-27.tsv', 'gsc_daily_', '.tsv'), '\t')
    _, d['gsc_insp'] = read_csv(find_file(data, 'gsc_inspection_2026-10-30.tsv', 'gsc_inspection_', '.tsv'), '\t')
    _, d['gsc_queries'] = read_csv(find_file(data, 'gsc_queries_v1_candidates.tsv', 'gsc_queries_', '.tsv'), '\t')
    readme = data / 'README.txt'
    d['readme_lines'] = [ln.strip() for ln in readme.read_text(encoding='utf-8').splitlines() if ln.strip()] \
        if readme.exists() else ['README.txt not found']

    # optional
    d['optional'] = {}
    for key, name in (('changelog', 'changelog.jsonl'), ('qa', 'qa_report.jsonl'),
                      ('checkpoint', 'checkpoint.json'), ('technical', 'technical_findings.json'),
                      ('repo_audit', 'repo_audit.json'), ('keyword_matrix', 'keyword_matrix.csv')):
        d['optional'][name] = (data / name).exists()
    d['changelog'] = read_jsonl(data / 'changelog.jsonl')
    d['qa'] = read_jsonl(data / 'qa_report.jsonl')
    d['checkpoint'] = read_json(data / 'checkpoint.json', None)
    d['technical'] = read_json(data / 'technical_findings.json', None) or []
    d['repo_audit'] = read_json(data / 'repo_audit.json', None) or []
    d['kwm_fields'], d['keyword_matrix'] = read_csv(data / 'keyword_matrix.csv')
    d['kw_metrics'] = read_jsonl(KW_METRICS)

    # indexes
    d['inv_by_url'] = {r['url']: r for r in d['inv']}
    d['cls_by_url'] = {r['url']: r for r in d['cls']}
    d['gsc_pages_by_url'] = {path_of(r.get('page')): r for r in d['gsc_pages']}
    d['insp_by_url'] = {path_of(r.get('url')): r for r in d['gsc_insp']}

    # keyword metrics: best record per keyword (prefer a volume, then the latest check)
    best = {}
    for r in d['kw_metrics']:
        k = (r.get('keyword') or '').strip().lower()
        if not k:
            continue
        score = (r.get('volume') is not None, str(r.get('checked') or ''))
        if k not in best or score > best[k][0]:
            best[k] = (score, r)
    d['kw_best'] = {k: v[1] for k, v in best.items()}

    # latest QA result per URL, latest changelog commit per URL
    qa_latest = {}
    for r in sorted(d['qa'], key=lambda x: str(x.get('ts') or '')):
        if r.get('url'):
            qa_latest[path_of(r['url'])] = r
    d['qa_latest'] = qa_latest
    commit_latest = {}
    for r in sorted(d['changelog'], key=lambda x: str(x.get('ts') or '')):
        if r.get('url') and r.get('commit'):
            commit_latest[path_of(r['url'])] = r['commit']
    d['commit_latest'] = commit_latest

    # hub -> hub page / pillar URL
    hub_pages, pillars = {}, {}
    for r in d['inv']:
        if r.get('page_type') == 'hub' and r.get('hub'):
            hub_pages[r['hub']] = r['url']
        elif r.get('page_type') == 'pillar' and r.get('hub'):
            pillars.setdefault(r['hub'], r['url'])
    d['hub_url'] = {h: hub_pages.get(h) or pillars.get(h) for h in set(hub_pages) | set(pillars)}
    return d


def qa_pass(r):
    v = r.get('pass', r.get('pass_'))
    if isinstance(v, str):
        return v.upper() if v.upper() in ('PASS', 'FAIL') else ('PASS' if is_true(v) else 'FAIL')
    return 'PASS' if v else 'FAIL'


# ----------------------------------------------------------------------------- sheets
def build_inventory(wb, d, R):
    sh = Sheet(wb, '02_ALL_URL_INVENTORY')
    headers = ['ID'] + d['inv_fields'] + ['category', 'action', 'status', 'priority', 'note']
    rows = []
    n = 0
    for r in d['inv']:
        n += 1
        c = d['cls_by_url'].get(r['url'], {})
        rows.append([f'U{n:05d}'] + [conv(col, r.get(col)) for col in d['inv_fields']]
                    + [clean(c.get('category')), clean(c.get('action')), clean(c.get('status')),
                       clean(c.get('priority')), None])
    for r in d['removed']:
        n += 1
        row = {'url': r.get('url'), 'abs_url': SITE + (r.get('url') or ''), 'page_type': 'REMOVED_PRIOR'}
        rows.append([f'U{n:05d}'] + [clean(row.get(col)) for col in d['inv_fields']]
                    + [None, None, clean(r.get('status')), None, clean(r.get('note'))])
    h, first, last = sh.table(headers, rows)
    sh.finish(headers, h, last, freeze='B2', wrap_cols=())
    sh.status_cf(sh.col(headers, 'status'), first, last)
    R['02'] = dict(headers=headers, first=first, last=last, sheet=sh)


def build_index_recovery(wb, d, R):
    sh = Sheet(wb, '03_INDEX_RECOVERY')
    headers = ['url', 'page_type', 'hub', 'noindex_before', 'noindex_reason', 'audit_decision', 'merge_into',
               'category', 'action', 'status', 'priority', 'reason', 'fix_method', 'noindex_after',
               'canonical_self', 'in_sitemap', 'qa_score', 'gsc_inspection', 'gsc_impressions',
               'gsc_clicks', 'commit', 'notes']
    rows = []
    for r in d['inv']:
        c = d['cls_by_url'].get(r['url'], {})
        if not (is_true(r.get('noindex')) or c.get('category') in RECOVERY_CATEGORIES or (c.get('status') or '')):
            continue
        qa = d['qa_latest'].get(r['url'])
        rows.append([clean(r['url']), clean(r.get('page_type')), clean(r.get('hub')), clean(r.get('noindex')),
                     clean(r.get('noindex_reason')), clean(c.get('audit_decision') or r.get('audit_decision')),
                     clean(c.get('merge_into') or r.get('merge_into')), clean(c.get('category')),
                     clean(c.get('action')), clean(c.get('status')), clean(c.get('priority')),
                     clean(c.get('reason')), None, None, clean(r.get('canonical_self')),
                     clean(r.get('in_sitemap')), num(qa.get('score')) if qa else None,
                     clean(c.get('gsc_inspection')), num(c.get('gsc_impressions')), num(c.get('gsc_clicks')),
                     clean(d['commit_latest'].get(r['url'])), None])
    h, first, last = sh.table(headers, rows)
    sh.finish(headers, h, last, freeze='B2')
    sh.status_cf(sh.col(headers, 'status'), first, last)
    sh.list_dv(sh.col(headers, 'status'), first, last, STATUS_VALUES)
    sh.list_dv(sh.col(headers, 'priority'), first, last, PRIORITY_VALUES)
    sh.list_dv(sh.col(headers, 'noindex_after'), first, last, ['TRUE', 'FALSE'])
    R['03'] = dict(headers=headers, first=first, last=last, sheet=sh)


def build_keyword_matrix(wb, d, R):
    sh = Sheet(wb, '04_KEYWORD_MATRIX')
    headers = ['Keyword', 'Intent', 'Cluster', 'Hub', 'Pillar/Hub URL', 'Target URL', 'Priority', 'Volume',
               'Volume source/date', 'KD', 'GSC impressions', 'GSC clicks', 'GSC position',
               'SEO Title (current)', 'Meta (current)', 'Internal links in (inlinks_total)', 'Status',
               'Index (gsc_inspection)', 'QA', 'Notes']
    kw_best, cls_by, inv_by, gsc_by = d['kw_best'], d['cls_by_url'], d['inv_by_url'], d['gsc_pages_by_url']

    def metrics(keyword):
        m = kw_best.get((keyword or '').strip().lower())
        if not m:
            return None, None, None, None
        src = f"{m.get('source') or ''} {m.get('checked') or ''}".strip() or None
        return num(m.get('volume')), src, num(m.get('kd')), (m.get('intent') or '').strip().lower() or None

    def qa_text(url):
        qa = d['qa_latest'].get(url)
        return f"{qa_pass(qa)} {qa.get('score', '')}".strip() if qa else None

    def page_fields(url):
        r, c, g = inv_by.get(url, {}), cls_by.get(url, {}), gsc_by.get(url, {})
        return r, c, g

    rows = []
    if d['keyword_matrix']:
        for k in d['keyword_matrix']:
            url = path_of(k.get('target_url'))
            r, c, g = page_fields(url)
            hub = k.get('hub') or r.get('hub')
            rows.append([clean(k.get('keyword')), clean(k.get('intent')), clean(k.get('cluster')), clean(hub),
                         clean(d['hub_url'].get(hub)), clean(url or None), clean(k.get('priority')),
                         num(k.get('volume')), clean(k.get('volume_source')), num(k.get('kd')),
                         num(k.get('gsc_impressions')), num(k.get('gsc_clicks')), num(k.get('gsc_position')),
                         clean(k.get('title') or r.get('title')), clean((k.get('meta') or r.get('meta_desc') or '')[:200]),
                         num(r.get('inlinks_total')), clean(k.get('status') or c.get('status')),
                         clean(c.get('gsc_inspection')), qa_text(url), clean(k.get('notes'))])
    else:
        for r in d['inv']:
            if is_true(r.get('noindex')) or r.get('page_type') in ('pagination', 'author'):
                continue
            c, g = cls_by.get(r['url'], {}), gsc_by.get(r['url'], {})
            kw, src = (r.get('primary_keyword') or '').strip(), 'primary_keyword'
            if not kw and (r.get('keywords') or '').strip():
                kw, src = r['keywords'].split(',')[0].strip(), 'keywords[0]'
            if not kw:
                kw, src = (r.get('title') or '').strip(), 'title'
            vol, vsrc, kd, intent = metrics(kw)
            notes = []
            if src != 'primary_keyword':
                notes.append(f'keyword fallback: {src}')
            if not intent:
                intent = 'commercial' if (r.get('hub') == 'thue-xe' or r.get('page_type') in ('pillar', 'home')) \
                    else 'informational'
                notes.append('intent: heuristic (no keyword-metrics record)')
            if vol is None:
                notes.append('no volume in data/keyword-metrics.jsonl')
            rows.append([clean(kw), intent, clean(c.get('cluster_id')), clean(r.get('hub')),
                         clean(d['hub_url'].get(r.get('hub'))), clean(r['url']), clean(c.get('priority')),
                         vol, clean(vsrc), kd, num(g.get('impressions') or c.get('gsc_impressions')),
                         num(g.get('clicks') or c.get('gsc_clicks')), num(g.get('position') or c.get('gsc_position')),
                         clean(r.get('title')), clean((r.get('meta_desc') or '')[:200]), num(r.get('inlinks_total')),
                         clean(c.get('status')), clean(c.get('gsc_inspection')), qa_text(r['url']),
                         clean('; '.join(notes))])
    for q in d['gsc_queries']:
        url = path_of(q.get('page'))
        r, c, g = page_fields(url)
        vol, vsrc, kd, intent = metrics(q.get('query'))
        rows.append([clean(q.get('query')), intent, clean(c.get('cluster_id')), clean(r.get('hub')),
                     clean(d['hub_url'].get(r.get('hub'))), clean(url), clean(c.get('priority')), vol, clean(vsrc), kd,
                     num(q.get('impressions')), num(q.get('clicks')), num(q.get('position')), clean(r.get('title')),
                     clean((r.get('meta_desc') or '')[:200]), num(r.get('inlinks_total')), clean(c.get('status')),
                     clean(c.get('gsc_inspection')), qa_text(url),
                     'GSC query candidate (gsc_queries_v1_candidates.tsv)' + ('' if r else '; page not in inventory')])
    h, first, last = sh.table(headers, rows)
    sh.finish(headers, h, last, freeze='B2')
    sh.status_cf(sh.col(headers, 'Status'), first, last)
    sh.number_format(sh.col(headers, 'GSC position'), first, last, '0.0')
    R['04'] = dict(headers=headers, first=first, last=last, sheet=sh)


def build_duplicates(wb, d, R):
    sh = Sheet(wb, '05_DUPLICATE_CLUSTERS')
    headers = ['cluster_id', 'cluster_size', 'url', 'hub', 'indexable', 'title', 'max_sim', 'max_sim_url',
               'max_sim_indexable', 'max_sim_indexable_url', 'audit_decision', 'merge_into', 'category',
               'action/decision', 'differentiation_plan', 'status']
    rows = []
    for s in d['sim']:
        if not (s.get('cluster_id') or to_f(s.get('max_sim_indexable')) >= 0.6):
            continue
        r, c = d['inv_by_url'].get(s['url'], {}), d['cls_by_url'].get(s['url'], {})
        rows.append([clean(s.get('cluster_id')), num(s.get('cluster_size')), clean(s['url']), clean(s.get('hub')),
                     clean(s.get('indexable')), clean(r.get('title') or c.get('title')),
                     round(to_f(s.get('max_sim')), 3), clean(s.get('max_sim_url')),
                     round(to_f(s.get('max_sim_indexable')), 3), clean(s.get('max_sim_indexable_url')),
                     clean(c.get('audit_decision') or r.get('audit_decision')),
                     clean(c.get('merge_into') or r.get('merge_into')), clean(c.get('category')),
                     clean(c.get('action')), None, clean(c.get('status'))])
    rows.sort(key=lambda x: (x[0] is None, x[0] or '', x[2] or ''))
    h, first, last = sh.table(headers, rows)
    sh.finish(headers, h, last, freeze='D2')
    sh.status_cf(sh.col(headers, 'status'), first, last)
    sh.list_dv(sh.col(headers, 'status'), first, last, STATUS_VALUES)
    sh.number_format(sh.col(headers, 'max_sim'), first, last, '0.000')
    sh.number_format(sh.col(headers, 'max_sim_indexable'), first, last, '0.000')
    R['05'] = dict(headers=headers, first=first, last=last, sheet=sh)


def content_flags(r):
    flags = []
    if to_i(r.get('title_len')) > 65:
        flags.append('TITLE>65')
    ml = to_i(r.get('meta_desc_len'))
    if ml > 160:
        flags.append('META>160')
    elif ml < 70:
        flags.append('META<70')
    if to_i(r.get('word_count')) < 900:
        flags.append('WORDS<900')
    if to_i(r.get('external_links')) == 0 and r.get('hub') in NEEDS_SOURCES:
        flags.append('NO_EXTERNAL_SOURCE')
    if to_i(r.get('inlinks_total')) <= 1:
        flags.append('FEW_INLINKS')
    if not is_true(r.get('h1_eq_title')):
        flags.append('H1_NE_TITLE')
    return '; '.join(flags) or None


def build_content_audit(wb, d, R):
    sh = Sheet(wb, '06_CONTENT_AUDIT')
    headers = ['url', 'page_type', 'hub', 'title', 'title_len', 'title_full_len', 'meta_desc_len', 'h1_eq_title',
               'h2_count', 'h3_count', 'word_count', 'has_faq', 'external_links', 'sources_count', 'img_count',
               'inlinks_total', 'flags', 'priority', 'action', 'factual_error', 'status']
    rows = []
    for r in d['inv']:
        c = d['cls_by_url'].get(r['url'], {})
        if is_true(r.get('noindex')) and c.get('action') != 'REWRITE_THEN_INDEX':
            continue
        rows.append([clean(r['url']), clean(r.get('page_type')), clean(r.get('hub')), clean(r.get('title')),
                     num(r.get('title_len')), num(r.get('title_full_len')), num(r.get('meta_desc_len')),
                     clean(r.get('h1_eq_title')), num(r.get('h2_count')), num(r.get('h3_count')),
                     num(r.get('word_count')), clean(r.get('has_faq')), num(r.get('external_links')),
                     num(r.get('sources_count')), num(r.get('img_count')), num(r.get('inlinks_total')),
                     content_flags(r), clean(c.get('priority')), clean(c.get('action')), None, clean(c.get('status'))])
    h, first, last = sh.table(headers, rows)
    sh.finish(headers, h, last, freeze='B2')
    sh.status_cf(sh.col(headers, 'status'), first, last)
    sh.list_dv(sh.col(headers, 'status'), first, last, STATUS_VALUES)
    sh.list_dv(sh.col(headers, 'priority'), first, last, PRIORITY_VALUES)
    R['06'] = dict(headers=headers, first=first, last=last, sheet=sh)


def technical_issues(r):
    issues = []
    if to_i(r.get('h1_count'), -1) != 1:
        issues.append('H1!=1')
    if not is_true(r.get('canonical_self')):
        issues.append('CANONICAL_NOT_SELF')
    noindex, in_sitemap = is_true(r.get('noindex')), bool((r.get('in_sitemap') or '').strip())
    if noindex and in_sitemap:
        issues.append('NOINDEX_IN_SITEMAP')
    if not noindex and not in_sitemap and r.get('page_type') != 'pagination':
        issues.append('INDEXABLE_NOT_IN_SITEMAP')
    if to_i(r.get('inlinks_total')) == 0 and r.get('page_type') in ARTICLE_TYPES:
        issues.append('ORPHAN')
    if to_i(r.get('img_missing_alt')) > 0:
        issues.append('IMG_ALT_MISSING')
    return '; '.join(issues) or None


def build_technical(wb, d, R):
    sh = Sheet(wb, '07_TECHNICAL_SEO')
    f_headers = ['area', 'finding', 'affected_count', 'evidence', 'fix', 'status', 'commit']
    f_rows = [[clean(x.get('area')), clean(x.get('finding')), num(x.get('affected_count')), joined(x.get('evidence')),
               joined(x.get('fix')), clean(x.get('status')), clean(x.get('commit'))]
              for x in d['technical'] if isinstance(x, dict)]
    fh, ffirst, flast = sh.table(f_headers, f_rows, title='A. Site-wide technical findings (technical_findings.json)')
    sh.blank(2)
    u_headers = ['url', 'page_type', 'robots', 'noindex', 'canonical_self', 'in_sitemap', 'h1_count',
                 'has_breadcrumb_schema', 'jsonld_types', 'img_missing_alt', 'og_image_default', 'inlinks_total',
                 'bytes', 'issues']
    u_rows = [[clean(r['url']), clean(r.get('page_type')), clean(r.get('robots')), clean(r.get('noindex')),
               clean(r.get('canonical_self')), clean(r.get('in_sitemap')), num(r.get('h1_count')),
               clean(r.get('has_breadcrumb_schema')), clean(r.get('jsonld_types')), num(r.get('img_missing_alt')),
               clean(r.get('og_image_default')), num(r.get('inlinks_total')), num(r.get('bytes')), technical_issues(r)]
              for r in d['inv']]
    uh, ufirst, ulast = sh.table(u_headers, u_rows, title='B. Per-URL technical checks (inventory.csv)')
    headers = u_headers if len(u_headers) >= len(f_headers) else f_headers
    freeze = f'B{uh + 1}' if uh <= 25 else 'A2'
    sh.finish(headers, uh, ulast, freeze=freeze, wrap_cols=())
    for name in WRAP_COLS:
        if name in f_headers:
            c = f_headers.index(name) + 1
            for r in range(ffirst, flast + 1):
                sh.ws.cell(r, c).alignment = WRAP
    sh.status_cf(sh.col(f_headers, 'status'), ffirst, flast)
    R['07'] = dict(headers=u_headers, first=ufirst, last=ulast, f_headers=f_headers, f_first=ffirst,
                   f_last=flast, sheet=sh)


def build_internal_links(wb, d, R):
    sh = Sheet(wb, '08_INTERNAL_LINKS')
    base = d['links_fields'] or ['source_url', 'target_url', 'link_type', 'source_indexable', 'target_exists',
                                 'target_noindex']
    headers = base + ['anchor_text', 'error', 'fix_status']
    rows = []
    for r in d['links']:
        err = []
        if r.get('target_exists') and not is_true(r.get('target_exists')):
            err.append('BROKEN')
        if is_true(r.get('target_noindex')):
            err.append('LINK_TO_NOINDEX')
        rows.append([clean(r.get(col)) for col in base] + [None, '; '.join(err) or None, None])
    h, first, last = sh.table(headers, rows)
    sh.finish(headers, h, last, freeze='A2', wrap_cols=())
    R['08'] = dict(headers=headers, first=first, last=last, sheet=sh)


def build_changelog(wb, d, R):
    sh = Sheet(wb, '09_CHANGELOG')
    headers = ['ts', 'batch', 'url', 'file', 'change_type', 'field', 'before', 'after', 'commit', 'note']
    rows = [[clean(x.get('ts')), clean(x.get('batch')), clean(x.get('url')), clean(x.get('file')),
             clean(x.get('change_type')), clean(x.get('field')), clean(x.get('before')), clean(x.get('after')),
             clean(x.get('commit')), clean(x.get('note'))] for x in d['changelog'] if isinstance(x, dict)]
    h, first, last = sh.table(headers, rows)
    sh.finish(headers, h, last)
    R['09'] = dict(headers=headers, first=first, last=last, sheet=sh)


def build_qa(wb, d, R):
    sh = Sheet(wb, '10_QA_REPORT')
    headers = ['ts', 'batch', 'url', 'score', 'pass', 'checks', 'critical', 'before', 'after',
               'production_verified', 'evidence']
    rows = []
    for x in d['qa']:
        if not isinstance(x, dict):
            continue
        checks = x.get('checks')
        rows.append([clean(x.get('ts')), clean(x.get('batch')), clean(x.get('url')), num(x.get('score')), qa_pass(x),
                     joined(checks) if isinstance(checks, dict) else clean(checks), joined(x.get('critical')),
                     clean(x.get('before')), clean(x.get('after')), clean(x.get('production_verified')),
                     joined(x.get('evidence'))])
    h, first, last = sh.table(headers, rows)
    sh.finish(headers, h, last)
    sh.status_cf(sh.col(headers, 'pass'), first, last)
    R['10'] = dict(headers=headers, first=first, last=last, sheet=sh)


def build_checkpoint(wb, d, R):
    sh = Sheet(wb, '11_CHECKPOINT')
    headers = ['key', 'value']
    rows = []
    cp = d['checkpoint']
    if isinstance(cp, dict):
        for key, val in cp.items():
            if isinstance(val, list):
                rows.append([f'{key}_count', len(val)])
                items = [v if isinstance(v, str) else json.dumps(v, ensure_ascii=False) for v in val]
                chunk, size, part = [], 0, 1
                for it in items:          # keep every cell under the Excel character limit
                    if chunk and size + len(it) + 1 > MAX_CELL:
                        rows.append([f'{key} (part {part})', clean('\n'.join(chunk))])
                        chunk, size, part = [], 0, part + 1
                    chunk.append(it)
                    size += len(it) + 1
                rows.append([key if part == 1 else f'{key} (part {part})', clean('\n'.join(chunk)) if chunk else None])
            elif isinstance(val, dict):
                rows.append([key, clean(json.dumps(val, ensure_ascii=False, indent=1))])
            else:
                rows.append([key, clean(val)])
    h, first, last = sh.table(headers, rows)
    sh.finish(headers, h, last, widths={'A': 28, 'B': 100})
    R['11'] = dict(headers=headers, first=first, last=last, sheet=sh)


def build_gsc(wb, d, R):
    sh = Sheet(wb, '12_GSC_PERFORMANCE')
    r = sh.append([GSC_NOTE])
    sh.ws.cell(r, 1).font = Font(bold=True, italic=True)
    sh.blank(1)
    info = {}

    daily_h = ['date', 'clicks', 'impressions']
    daily_rows = [[clean(x.get('date')), num(x.get('clicks')), num(x.get('impressions'))] for x in d['gsc_daily']]
    h, first, last = sh.table(daily_h, daily_rows, title='A. Daily clicks / impressions (gsc_daily)')
    info['daily'] = (h, first, last)
    sh.blank(2)

    pages_h = ['page', 'clicks', 'impressions', 'ctr', 'position', 'indexing_status']
    pages_rows = []
    for x in d['gsc_pages']:
        p = path_of(x.get('page'))
        insp = d['insp_by_url'].get(p)
        pages_rows.append([clean(p), num(x.get('clicks')), num(x.get('impressions')), num(x.get('ctr')),
                           num(x.get('position')), clean(insp.get('coverageState')) if insp else 'NOT_VERIFIED'])
    h, first, last = sh.table(pages_h, pages_rows, title='B. Pages (gsc_pages) + indexing status from URL Inspection')
    info['pages'] = (h, first, last)
    sh.number_format('D', first, last, '0.0%')
    sh.number_format('E', first, last, '0.0')
    sh.blank(2)

    insp_h = ['url', 'coverageState', 'verdict', 'lastCrawlTime', 'googleCanonical_self', 'note']
    insp_rows = [[clean(path_of(x.get('url'))), clean(x.get('coverageState')), clean(x.get('verdict')),
                  clean(x.get('lastCrawlTime')), clean(x.get('googleCanonical_self')), clean(x.get('note'))]
                 for x in d['gsc_insp']]
    h, first, last = sh.table(insp_h, insp_rows, title='C. URL Inspection (gsc_inspection)')
    info['insp'] = (h, first, last)
    sh.blank(2)

    q_h = ['page', 'query', 'clicks', 'impressions', 'position']
    q_rows = [[clean(path_of(x.get('page'))), clean(x.get('query')), num(x.get('clicks')), num(x.get('impressions')),
               num(x.get('position'))] for x in d['gsc_queries']]
    h, first, last = sh.table(q_h, q_rows, title='D. Query candidates (gsc_queries_v1_candidates)')
    info['queries'] = (h, first, last)
    sh.number_format('E', first, last, '0.0')

    ph, pfirst, plast = info['pages']
    sh.finish(pages_h, ph, plast, freeze='A2', wrap_cols=())
    sh.ws.column_dimensions['A'].width = 60
    R['12'] = dict(info=info, sheet=sh)


def build_repo_audit(wb, d, R):
    sh = Sheet(wb, '13_REPOSITORY_AUDIT')
    headers = ['item', 'value', 'risk', 'note']
    rows = [[clean(x.get('item')), clean(x.get('value')), clean(x.get('risk')), clean(x.get('note'))]
            for x in d['repo_audit'] if isinstance(x, dict)]
    h, first, last = sh.table(headers, rows)
    sh.finish(headers, h, last)
    R['13'] = dict(headers=headers, first=first, last=last, sheet=sh)


def build_dashboard(wb, d, R, data_dir, generated):
    sh = Sheet(wb, '01_DASHBOARD')
    ws = sh.ws
    hdr = sh.header(['Metric', 'Value', 'Note'])
    sh.append(['Last generated', generated, 'Python datetime at run (local time); re-run scripts/seo_matrix_export.py to refresh'])
    sh.append(['Site', SITE, ''])
    sh.append(['Data directory', str(data_dir), ''])
    for i, line in enumerate(d['readme_lines'], 1):
        sh.append([f'Data provenance {i}', line, 'README.txt'])

    def section(title):
        r = sh.append([title])
        for c in (1, 2, 3):
            ws.cell(r, c).fill = SECTION_FILL
        ws.cell(r, 1).font = HEADER_FONT

    def kpi(label, value, note='', fmt=None):
        r = sh.append([label, value, note])
        if fmt:
            ws.cell(r, 2).number_format = fmt
        return f'B{r}'

    def colref(key, name):
        # whole column below the header (row 2..max) so header text never matches a COUNTIF criterion
        s = R[key]
        L = get_column_letter(s['headers'].index(name) + 1)
        return f"'{SHEET_NAMES[int(key) - 1]}'!{L}2:{L}1048576"

    def rng(key, name, first, last):
        s = R[key]
        L = get_column_letter(s['headers'].index(name) + 1)
        return f"'{SHEET_NAMES[int(key) - 1]}'!{L}{first}:{L}{last}"

    def distinct(rows, col):
        return sorted({(r.get(col) or '') for r in rows if (r.get(col) or '')})

    s02, s03 = SHEET_NAMES[1], SHEET_NAMES[2]
    pt, nx, nr, cat, sm = (colref('02', c) for c in ('page_type', 'noindex', 'noindex_reason', 'category', 'in_sitemap'))
    tfl, mdl = colref('02', 'title_full_len'), colref('02', 'meta_desc_len')

    sh.blank()
    section('A. URL inventory (02_ALL_URL_INVENTORY)')
    total = kpi('Rows in 02 (incl. REMOVED_PRIOR)', f"=COUNTA('{s02}'!A:A)-1", 'COUNTA of ID column minus header')
    removed = kpi('REMOVED_PRIOR rows (deleted before this project)', f'=COUNTIF({pt},"REMOVED_PRIOR")',
                  'removed_prior_*.csv')
    kpi('Rendered URLs in inventory', f'={total}-{removed}', 'inventory.csv')
    for t in PAGE_TYPES:
        kpi(f'page_type = {t}', f'=COUNTIF({pt},"{t}")')
    kpi('noindex = True', f'=COUNTIF({nx},"True")', 'robots meta noindex at audit time')
    kpi('noindex = False (index-eligible)', f'=COUNTIF({nx},"False")')
    kpi('Index-eligible excl. pagination/author',
        f'=COUNTIFS({nx},"False",{pt},"<>pagination",{pt},"<>author")')
    for reason in distinct(d['inv'], 'noindex_reason'):
        kpi(f'noindex_reason = {reason}', f'=COUNTIF({nr},"{reason}")')
    kpi('URLs in a sitemap', f'=COUNTIF({sm},"sitemap*")', 'in_sitemap non-empty')
    for smf in distinct(d['inv'], 'in_sitemap'):
        kpi(f'in_sitemap = {smf}', f'=COUNTIF({sm},"{smf}")')
    kpi('title_full_len > 65 among noindex=False', f'=COUNTIFS({nx},"False",{tfl},">65")',
        'full <title> incl. brand suffix')
    kpi('meta_desc_len > 160 among noindex=False', f'=COUNTIFS({nx},"False",{mdl},">160")')
    kpi('meta_desc_len < 70 among noindex=False', f'=COUNTIFS({nx},"False",{mdl},"<70")')

    sh.blank()
    section('B. Classification & recovery progress (category from 02; action/status/priority from 03_INDEX_RECOVERY)')
    for c in CATEGORIES:
        kpi(f'category = {c}', f'=COUNTIF({cat},"{c}")', 'classification.csv')
    act, st, pr = colref('03', 'action'), colref('03', 'status'), colref('03', 'priority')
    kpi('Rows in 03_INDEX_RECOVERY', f"=COUNTA('{s03}'!A:A)-1")
    for a in distinct(d['cls'], 'action'):
        kpi(f'action = {a}', f'=COUNTIF({act},"{a}")')
    st_cells = {}
    for s in STATUS_VALUES:
        st_cells[s] = kpi(f'status = {s}', f'=COUNTIF({st},"{s}")', 'editable in 03')
    for p in PRIORITY_VALUES:
        kpi(f'priority = {p}', f'=COUNTIF({pr},"{p}")')
    kpi('Progress % = DONE / (OPEN + DONE + IN_PROGRESS)',
        f"=IFERROR({st_cells['DONE']}/({st_cells['OPEN']}+{st_cells['DONE']}+{st_cells['IN_PROGRESS']}),0)",
        'sheet 03 status column', fmt='0.0%')
    kpi('noindex_after = FALSE recorded in 03', f'=COUNTIF({colref("03", "noindex_after")},"FALSE")',
        'URLs re-opened to index (editable)')
    kpi('qa_score filled in 03', f'=COUNT({colref("03", "qa_score")})', 'latest qa_report.jsonl score per URL')

    sh.blank()
    section('C. Content, duplicates, technical, internal links')
    kpi('06 rows (indexable + REWRITE_THEN_INDEX)', f"=COUNTA('{SHEET_NAMES[5]}'!A:A)-1")
    fl = colref('06', 'flags')
    kpi('06 rows with any flag', f'=COUNTIF({fl},"?*")')
    for f in ('TITLE>65', 'META>160', 'META<70', 'WORDS<900', 'NO_EXTERNAL_SOURCE', 'FEW_INLINKS', 'H1_NE_TITLE'):
        kpi(f'06 flag {f}', f'=COUNTIF({fl},"*{f}*")')
    kpi('05 rows (cluster members / near-duplicates)', f"=COUNTA('{SHEET_NAMES[4]}'!C:C)-1",
        'cluster_id non-empty OR max_sim_indexable >= 0.6')
    r05 = R['05']
    if r05['last'] >= r05['first']:
        cid = rng('05', 'cluster_id', r05['first'], r05['last'])
        kpi('05 distinct cluster_id', f'=SUMPRODUCT(({cid}<>"")/COUNTIF({cid},{cid}&""))', 'explicit range')
        kpi('05 rows with max_sim_indexable >= 0.6',
            f'=COUNTIF({rng("05", "max_sim_indexable", r05["first"], r05["last"])},">=0.6")')
    r07 = R['07']
    iss = rng('07', 'issues', r07['first'], r07['last'])
    kpi('07 per-URL rows with any issue', f'=COUNTIF({iss},"?*")', 'block B explicit range')
    for i in ('H1!=1', 'CANONICAL_NOT_SELF', 'NOINDEX_IN_SITEMAP', 'INDEXABLE_NOT_IN_SITEMAP', 'ORPHAN',
              'IMG_ALT_MISSING'):
        kpi(f'07 issue {i}', f'=COUNTIF({iss},"*{i}*")')
    kpi('07 ORPHAN among noindex=False (index-eligible)',
        f'=COUNTIFS({iss},"*ORPHAN*",{rng("07", "noindex", r07["first"], r07["last"])},"False")')
    kpi('07 site-wide findings', f"=COUNTA('{SHEET_NAMES[6]}'!A{r07['f_first']}:A{max(r07['f_last'], r07['f_first'])})",
        'technical_findings.json' + ('' if d['optional']['technical_findings.json'] else ' (absent: header only)'))
    kpi('08 internal link edges', f"=COUNTA('{SHEET_NAMES[7]}'!A:A)-1", 'internal_links.csv')
    err = colref('08', 'error')
    kpi('08 BROKEN links', f'=COUNTIF({err},"*BROKEN*")')
    kpi('08 links to noindex pages', f'=COUNTIF({err},"*LINK_TO_NOINDEX*")')
    kpi('08 links fixed (fix_status = DONE)', f'=COUNTIF({colref("08", "fix_status")},"DONE")', 'editable')

    sh.blank()
    section('D. Work log (09_CHANGELOG, 10_QA_REPORT, 11_CHECKPOINT)')
    kpi('Changelog entries', f"=COUNTA('{SHEET_NAMES[8]}'!A:A)-1",
        'changelog.jsonl' + ('' if d['optional']['changelog.jsonl'] else ' (absent: header only)'))
    kpi('Changelog distinct commits', f'=SUMPRODUCT(({colref("09","commit")}<>"")/COUNTIF({colref("09","commit")},{colref("09","commit")}&""))'
        if R['09']['last'] >= R['09']['first'] else '=0', 'whole column; 0 when empty')
    kpi('QA entries', f"=COUNTA('{SHEET_NAMES[9]}'!A:A)-1",
        'qa_report.jsonl' + ('' if d['optional']['qa_report.jsonl'] else ' (absent: header only)'))
    kpi('QA PASS', f'=COUNTIF({colref("10", "pass")},"PASS")')
    kpi('QA FAIL', f'=COUNTIF({colref("10", "pass")},"FAIL")')
    kpi('QA production_verified = True', f'=COUNTIF({colref("10", "production_verified")},"True")')
    cp = d['checkpoint'] if isinstance(d['checkpoint'], dict) else {}
    kpi('Checkpoint batch_id', clean(cp.get('batch_id')) or 'n/a',
        'checkpoint.json' + ('' if d['optional']['checkpoint.json'] else ' (absent)'))
    kpi('Checkpoint commit_sha', clean(cp.get('commit_sha')) or 'n/a')
    kpi('Checkpoint next_task', clean(cp.get('next_task')) or 'n/a')
    kpi('Checkpoint updated', clean(cp.get('updated')) or 'n/a')
    cu = cp.get('completed_urls')
    kpi('Checkpoint completed_urls', len(cu) if isinstance(cu, list) else (clean(cu) or 0))
    ru = cp.get('remaining_urls')
    kpi('Checkpoint remaining_urls', len(ru) if isinstance(ru, list) else (num(ru) if ru is not None else 'n/a'))
    oe = cp.get('open_errors')
    kpi('Checkpoint open_errors', len(oe) if isinstance(oe, list) else (clean(oe) or 0))

    sh.blank()
    section('E. Google Search Console (12_GSC_PERFORMANCE)')
    g = R['12']['info']
    s12 = SHEET_NAMES[11]
    dh, df, dl = g['daily']
    if dl >= df:
        kpi('Daily block: total clicks', f"=SUM('{s12}'!B{df}:B{dl})", f'rows {df}-{dl}')
        kpi('Daily block: total impressions', f"=SUM('{s12}'!C{df}:C{dl})")
        kpi('Daily block: days', f"=COUNTA('{s12}'!A{df}:A{dl})")
        kpi('Daily block: first date', f"='{s12}'!A{df}")
        kpi('Daily block: last date', f"='{s12}'!A{dl}")
    ph, pf, pl = g['pages']
    if pl >= pf:
        kpi('Pages block: pages with data', f"=COUNTA('{s12}'!A{pf}:A{pl})")
        kpi('Pages block: total clicks', f"=SUM('{s12}'!B{pf}:B{pl})")
        kpi('Pages block: total impressions', f"=SUM('{s12}'!C{pf}:C{pl})")
        kpi('Pages block: pages with clicks > 0', f"=COUNTIF('{s12}'!B{pf}:B{pl},\">0\")")
        kpi('Pages block: avg position (impression-weighted)',
            f"=IFERROR(SUMPRODUCT('{s12}'!C{pf}:C{pl},'{s12}'!E{pf}:E{pl})/SUM('{s12}'!C{pf}:C{pl}),0)", '', fmt='0.0')
        kpi('Pages block: NOT_VERIFIED by inspection', f"=COUNTIF('{s12}'!F{pf}:F{pl},\"NOT_VERIFIED\")")
    ih, if_, il = g['insp']
    if il >= if_:
        kpi('Inspection block: URLs inspected', f"=COUNTA('{s12}'!A{if_}:A{il})")
        for state in ('Submitted and indexed', 'Crawled - currently not indexed',
                      'Discovered - currently not indexed', 'URL is unknown to Google'):
            kpi(f'Inspection: {state}', f"=COUNTIF('{s12}'!B{if_}:B{il},\"{state}\")")
        kpi('Inspection: verdict PASS', f"=COUNTIF('{s12}'!C{if_}:C{il},\"PASS\")")
    qh, qf, ql = g['queries']
    kpi('Query candidates (block D)', f"=COUNTA('{s12}'!A{qf}:A{max(ql, qf - 1)})" if ql >= qf else '=0')
    kpi('04_KEYWORD_MATRIX rows', f"=COUNTA('{SHEET_NAMES[3]}'!A:A)-1",
        'keyword_matrix.csv' if d['optional']['keyword_matrix.csv'] else
        'built from classification/inventory + data/keyword-metrics.jsonl + GSC queries')
    kpi('04 rows with search volume', f'=COUNT({colref("04", "Volume")})')

    sh.blank()
    section('F. Input files')
    for name, present in d['optional'].items():
        kpi(f'optional: {name}', 'present' if present else 'missing (sheet has header only)')
    kpi('data/keyword-metrics.jsonl records', len(d['kw_metrics']), str(KW_METRICS))
    kpi('inventory.csv rows', len(d['inv']))
    kpi('classification.csv rows', len(d['cls']))
    kpi('similarity.csv rows', len(d['sim']))
    kpi('internal_links.csv rows', len(d['links']))
    kpi('removed_prior rows', len(d['removed']))
    kpi('gsc_pages / gsc_daily / gsc_inspection / gsc_queries rows',
        f"{len(d['gsc_pages'])} / {len(d['gsc_daily'])} / {len(d['gsc_insp'])} / {len(d['gsc_queries'])}")

    sh.finish(['Metric', 'Value', 'Note'], hdr, sh.row, freeze='A2', wrap_cols=(),
              widths={'A': 58, 'B': 44, 'C': 70})
    ws.cell(1, 1).value = 'Metric'


# ----------------------------------------------------------------------------- main
def export(data_dir, out_path):
    t0 = time.time()
    data_dir, out_path = Path(data_dir), Path(out_path)
    d = load_all(data_dir)
    if not d['inv']:
        sys.exit(f'ERROR: {data_dir / "inventory.csv"} missing or empty')
    print(f'loaded data in {time.time() - t0:.1f}s: inventory={len(d["inv"])} classification={len(d["cls"])} '
          f'similarity={len(d["sim"])} links={len(d["links"])} removed={len(d["removed"])} '
          f'kw_metrics={len(d["kw_metrics"])} optional={ {k: v for k, v in d["optional"].items()} }')
    wb = Workbook()
    wb.active.title = SHEET_NAMES[0]
    R = {}
    build_inventory(wb, d, R)
    build_index_recovery(wb, d, R)
    build_keyword_matrix(wb, d, R)
    build_duplicates(wb, d, R)
    build_content_audit(wb, d, R)
    build_technical(wb, d, R)
    build_internal_links(wb, d, R)
    build_changelog(wb, d, R)
    build_qa(wb, d, R)
    build_checkpoint(wb, d, R)
    build_gsc(wb, d, R)
    build_repo_audit(wb, d, R)
    generated = datetime.datetime.now().astimezone().isoformat(timespec='seconds')
    build_dashboard(wb, d, R, data_dir, generated)
    wb._sheets = [wb[n] for n in SHEET_NAMES]
    wb.active = 0
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f'wrote {out_path} ({out_path.stat().st_size / 1e6:.1f} MB) in {time.time() - t0:.1f}s')
    return out_path


def verify(out_path):
    wb = load_workbook(out_path, read_only=True)
    names = wb.sheetnames
    print('sheets:', names)
    ok = names == SHEET_NAMES
    for ws in wb.worksheets:
        print(f'  {ws.title:24s} rows={ws.max_row:6d} cols={ws.max_column}')
    wb.close()
    wb = load_workbook(out_path, read_only=False)
    ws = wb[SHEET_NAMES[0]]
    refs, bad, n_formulas = set(), set(), 0
    for row in ws.iter_rows(min_col=2, max_col=2):
        v = row[0].value
        if isinstance(v, str) and v.startswith('='):
            n_formulas += 1
            for m in re.finditer(r"'([^']+)'!", v):
                refs.add(m.group(1))
                if m.group(1) not in names:
                    bad.add(m.group(1))
    print(f'dashboard formulas: {n_formulas}; sheets referenced: {sorted(refs)}; unknown refs: {sorted(bad)}')
    print('VERIFY', 'OK' if ok and not bad else 'FAILED')
    return ok and not bad


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--data', default=str(DEFAULT_DATA), help=f'data directory (default: {DEFAULT_DATA})')
    ap.add_argument('--out', default=str(DEFAULT_OUT), help=f'output .xlsx (default: {DEFAULT_OUT})')
    ap.add_argument('--verify', action='store_true', help='reopen the workbook and print sheet names/row counts')
    args = ap.parse_args()
    out = export(args.data, args.out)
    if args.verify:
        sys.exit(0 if verify(out) else 1)


if __name__ == '__main__':
    main()
